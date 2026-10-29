#!/usr/bin/env python3
"""Build the five-page course report using the teacher's DOCX template.

Usage: python3 scripts/build_report.py --template /path/to/template.docx
Requires python-docx; PDF export requires LibreOffice (soffice) and pypdf.
On macOS, a temporary Fontconfig alias renders SimSun as Songti SC and
embeds the local Songti/Times fonts without changing system configuration.
The original cover and assessment tables are retained; text uses SimSun and
Times New Roman, 10.5 pt, 1.25 line spacing. No scan results are manufactured.
"""
from __future__ import annotations
import argparse
from copy import deepcopy
from pathlib import Path
import json
import os
import shutil
import subprocess
import tempfile
from datetime import datetime, timezone
import xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
STEM = 'F250117-F2505437-宋佳威-许可证测试作业'
TITLE = '许可证测试作业：Apache Commons CLI 合规分析'
PAGE_PENDING = False

def style_run(r, size=10.5, bold=False):
    r.font.name = 'Times New Roman'
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor(0, 0, 0)
    fonts = r._element.get_or_add_rPr().get_or_add_rFonts()
    for key, val in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia','宋体'),('cs','Times New Roman')]:
        fonts.set(qn('w:'+key), val)

def style_para(p, indent=True, size=10.5, bold=False, align=None, after=4):
    f=p.paragraph_format
    f.line_spacing = 1.25
    f.first_line_indent = Pt(21 if indent else 0)
    f.space_before = Pt(0)
    f.space_after = Pt(after)
    f.widow_control = True
    snap = p._p.get_or_add_pPr().find(qn('w:snapToGrid'))
    if snap is None:
        snap = OxmlElement('w:snapToGrid')
        p._p.get_or_add_pPr().append(snap)
    snap.set(qn('w:val'), '0')
    if align is not None: p.alignment=align
    for r in p.runs: style_run(r,size,bold)
    return p

def add_p(d,text,**kwargs):
    global PAGE_PENDING
    p = style_para(d.add_paragraph(text),**kwargs)
    if PAGE_PENDING:
        p.paragraph_format.page_break_before = True
        PAGE_PENDING = False
    return p

def heading(d,text,level=1):
    p=add_p(d,text,indent=False,size=12 if level==1 else 10.5,bold=True,after=6 if level==1 else 3)
    p.paragraph_format.keep_with_next=True
    return p

def page_break(d):
    global PAGE_PENDING
    PAGE_PENDING = True

def fix_table(t, size=10.5, center=False):
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for row in t.rows:
        trPr=row._tr.get_or_add_trPr()
        if trPr.find(qn('w:cantSplit')) is None: trPr.append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                style_para(p,indent=False,size=size,after=2,align=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT)
                p.paragraph_format.space_before=Pt(2)

def table(d, headers, rows, widths=None):
    t=d.add_table(rows=1, cols=len(headers))
    t.style='Table Grid'
    for c,x in zip(t.rows[0].cells,headers): c.text=x
    for row in rows:
        for c,x in zip(t.add_row().cells,row): c.text=str(x)
    fix_table(t)
    for p in t.rows[0].cells:
        for r in p.paragraphs[0].runs: r.bold=True
    if widths:
        t.autofit=False
        for row in t.rows:
            for c,w in zip(row.cells,widths): c.width=Pt(w)
    return t

def link_p(d,label,url):
    p=add_p(d,label+' ',indent=False,after=3)
    rid=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),rid)
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),'Times New Roman');fonts.set(qn('w:hAnsi'),'Times New Roman');fonts.set(qn('w:eastAsia'),'宋体');pr.append(fonts)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),'21');pr.append(sz)
    u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u)
    r.append(pr); tx=OxmlElement('w:t');tx.text=url; r.append(tx);h.append(r);p._p.append(h)
    return p

def build(template:Path, export_pdf=True):
    proof=json.loads((ROOT/'evidence/provenance.json').read_text())
    xml=ET.parse(ROOT/'evidence/rat-report.xml').getroot()
    resources=list(xml)
    counts={key:sum(x.find('type').get('name')==key for x in resources) for key in ('standard','notice','binary')}
    assert len(resources)==116 and counts=={'standard':109,'notice':3,'binary':4}
    approved=[x for x in resources if x.find('license-approval') is not None and x.find('license-approval').get('name')=='true']
    unknown=[x.get('name') for x in resources if x.find('license-approval') is not None and x.find('license-approval').get('name')=='false']
    assert len(approved)==106 and len(unknown)==3
    assert proof['project']['commit']=='698b238276c0e22e97e4aec703a0b00201d29666'
    d=Document(template)
    cover_xml=deepcopy(d.tables[0]._element)
    eval_xml=deepcopy(d.tables[1]._element)
    final_sect=deepcopy(d.sections[-1]._sectPr)
    for e in list(d._element.body): d._element.body.remove(e)
    d._element.body.append(final_sect)
    for s in d.sections:
        grid = s._sectPr.find(qn('w:docGrid'))
        if grid is not None:
            s._sectPr.remove(grid)
        s.top_margin=Pt(58);s.bottom_margin=Pt(52);s.left_margin=Pt(72);s.right_margin=Pt(72)
        s.header_distance=Pt(20);s.footer_distance=Pt(22)
        for p in s.header.paragraphs: p.clear()
        for p in s.footer.paragraphs: p.clear()
        footer=s.footer.paragraphs[0]
        style_para(footer,indent=False,size=9,align=WD_ALIGN_PARAGRAPH.CENTER,after=0)
        for txt in ['第 ','PAGE',' 页 / 共 ','NUMPAGES',' 页']:
            if txt in ('PAGE','NUMPAGES'):
                field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),txt);footer._p.append(field)
            else: style_run(footer.add_run(txt),9)
    normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(10.5)
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
    normal.paragraph_format.line_spacing=1.25
    d.core_properties.title=TITLE
    d.core_properties.author='宋佳威'
    d.core_properties.last_modified_by='宋佳威（AI辅助）'
    build_time=datetime.now(timezone.utc)
    d.core_properties.created=build_time
    d.core_properties.modified=build_time
    d.core_properties.subject='《开源创新思维》平时作业1：第2章开源合规与许可证'
    d.core_properties.keywords='F250117,F2505437,Apache Commons CLI,Apache Rat,Apache-2.0'
    # Cover: original template table, with its merged cells and assessment fields.
    add_p(d,'课程编号：',indent=False,size=10.5,after=30)
    add_p(d,'《开源创新思维》',indent=False,size=24,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,after=12)
    add_p(d,'平时作业报告',indent=False,size=24,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,after=24)
    add_p(d,TITLE,indent=False,size=14,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,after=6)
    add_p(d,'作业1 · 第2章 开源合规与许可证',indent=False,size=12,align=WD_ALIGN_PARAGRAPH.CENTER,after=28)
    d._element.body.insert(len(d._element.body)-1,cover_xml)
    t=d.tables[-1]
    t.cell(0,1).text='宋佳威';t.cell(0,4).text='F2505437'
    t.cell(1,1).text='F250117';t.cell(1,4).text='祝欣蓉'
    t.cell(2,1).text='2026-2027秋季学期'
    t.cell(3,1).text='第1周——第16周'
    t.cell(4,1).text='2026年10月29日'
    t.cell(5,1).text='';t.cell(5,4).text='祝欣蓉';t.cell(6,4).text=''
    fix_table(t,size=12,center=True)
    add_p(d,'年级：大二',indent=False,size=12,align=WD_ALIGN_PARAGRAPH.CENTER,after=16)
    add_p(d,'上海杉达学院',indent=False,size=18,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,after=15)
    add_p(d,'报告与完整原始证据随仓库提交；扫描对象为 Apache Commons CLI。',indent=False,size=10.5,align=WD_ALIGN_PARAGRAPH.CENTER,after=0)
    page_break(d)
    # Page 2
    heading(d,'1、开源项目介绍')
    add_p(d,'本次选择 Apache 软件基金会 Apache Commons 项目下的 Commons CLI。它是 Java 命令行参数解析库，可定义选项、解析参数并生成帮助信息。例如，程序可用它区分“--help”和带值的参数。仓库为 apache/commons-cli，项目许可证为 Apache-2.0。[1][2]')
    add_p(d,'为保证复现，选择固定版本 1.9.0，标签 rel/commons-cli-1.9.0；未将其描述为最新版本。源码提交为：',after=2)
    add_p(d,proof['project']['commit'],indent=False,after=5)
    add_p(d,'被测输入是 GitHub 提交源码快照，并非 Apache 正式发行压缩包。扫描解压后的全部 116 个文件，包括 .github 等隐藏目录，无自定义排除。范围涵盖 Java 主源码、测试、配置、网站资源和许可文件，不含 Git 历史、Maven 下载依赖、构建产物或容器镜像。')
    heading(d,'2、工具选择、安装过程及问题')
    add_p(d,'选用 Apache Creadur 的 Apache Rat 0.16.1。它适合检查 Apache 源码的许可证头及发布材料；工具本身也采用 Apache-2.0。相较于可识别更多许可证、版权和包信息的 ScanCode Toolkit，Rat 检查范围更窄；FOSSology 提供多扫描器和协作复核，但部署更重。本实验聚焦源码发布前的许可证头检查，因此采用易于复跑的 Rat 单个 JAR，并保留全部原始结果。[3]')
    add_p(d,'实验环境为 macOS 26.6.2（arm64）和已安装的 OpenJDK 17.0.18。先运行 java -version 确认环境；从 Maven Central 获取官方 org.apache.rat:apache-rat:0.16.1 JAR，执行 --help 验证可运行；从 GitHub 下载上述固定提交的源码压缩包并解压到 source。此方式不需要安装 Maven，也没有重新安装 Java。')
    add_p(d,'下载时沙箱网络存在 DNS/连接限制，取得网络执行授权后完成下载。环境、下载及扫描日志分别见 evidence/environment.txt、installation.log 和 scan-*.log；未把未发生的安装失败写作实验经历。工具及源码下载地址、SHA-256 均记入 provenance.json；本地散列用于固定实验输入，不替代发布者签名验证。')
    add_p(d,'核心扫描命令如下；完整下载、核验和复跑步骤见同目录 README.md。',after=3)
    add_p(d,'java -jar tools/apache-rat-0.16.1.jar --scan-hidden-directories source > rat-report.txt',indent=False,after=3)
    add_p(d,'java -jar tools/apache-rat-0.16.1.jar --scan-hidden-directories -x source > rat-report.xml',indent=False,after=3)
    add_p(d,'首次记录时间为 2026-10-29 13:29:09（北京时间）。只使用报告参数，不使用自动加许可证头参数，避免改变被测源码。',after=0)
    page_break(d)
    # Page 3
    heading(d,'3、工具输出的检测报告')
    add_p(d,'以下为 evidence/rat-report.txt 的原始摘要与待核文件清单摘录。完整逐文件 TXT 和原始 XML 均随仓库提交；为满足五页限制，报告正文不重复粘贴整份 116 文件清单。')
    raw='''Summary
-------
Generated at: 2026-10-29T13:29:09+08:00

Notes: 3
Binaries: 4
Archives: 0
Standards: 109

Apache Licensed: 106
Generated Documents: 0

3 Unknown Licenses

Files with unapproved licenses:
  source/.gitignore
  source/src/site/resources/download_cli.cgi
  source/src/site/resources/profile.jacoco'''
    for line in raw.splitlines():
        p=add_p(d,line or ' ',indent=False,after=0)
        p.paragraph_format.line_spacing=1.25
    add_p(d,'',indent=False,after=3)
    table(d,['统计项','数量','判读'],[
      ('标准文本资源','109','106 个 AL + 3 个 Unknown'),
      ('通知/许可类资源','3','LICENSE、NOTICE、RELEASE-NOTES'),
      ('二进制资源','4','logo.xcf、1x1.gif、logo.png、options.png'),
      ('全部资源','116','109 + 3 + 4；与源文件清单一致')],widths=[104,42,305])
    add_p(d,'AL 表示识别到 Apache License 2.0 文件头；N/B 分别表示通知类/二进制文件。106/109≈97.25% 仅为标准文本资源的头部识别比例，不是“项目合规率”。3 份通知类文件和 4 份二进制文件没有经过同样的许可证头批准判断，不能并入已批准数。',after=0)
    page_break(d)
    # Page 4
    heading(d,'4、检测结果详细分析')
    heading(d,'4.1 主要证据与适用范围',2)
    add_p(d,'XML 中 106 个标准文件标记为 Apache License Version 2.0 且 license-approval=true。主源码 26 个 Java 文件、测试目录 39 个 Java 文件均为 AL；LICENSE.txt 包含 Apache-2.0 正文，NOTICE.txt 载明 Apache Commons CLI、2002—2024 年 ASF 版权及基金会归属声明。这些证据支持“所扫主要 Java 源码许可证头一致”的结论，不能证明所有法律义务均已满足。')
    heading(d,'4.2 三项 Unknown 的逐项人工复核',2)
    add_p(d,'（1）.gitignore：内容为构建产物及 IDE 忽略规则，开头包含指向 github/gitignore 固定提交的来源 URL，但没有可识别许可证头。应核实对应提交的来源及适用许可，并在确有授权时补充适当说明；若发布审查允许排除此类配置，应记录理由、范围和复核人。当前保留 Unknown，不根据仓库总许可证擅自给外部来源内容改授权。')
    add_p(d,'（2）src/site/resources/download_cli.cgi：是带 #!/bin/sh 的短脚本，调用 Apache 的 mirrors.cgi，没有可识别许可证头。应核查提交历史及版权来源，确认授权后在 shebang 后加入合法许可证注释并复扫；若仅供网站部署、不进入发行包，也应验证发布清单并记录处置。没有来源证据前不能直接自动补 Apache 头。')
    add_p(d,'（3）src/site/resources/profile.jacoco：实际为 0 字节，Rat 无文本可匹配。空文件不等于许可证违规；应检查构建/站点是否需要它。若为必要占位文件，记录“空文件、无可扫描文本”的经复核例外；若无用途可在维护分支删除后验证构建。此次不修改上游快照，原始结果保留待核状态。')
    heading(d,'4.3 其他风险和检查充分性',2)
    add_p(d,'4 个图像/设计文件只被分类为 binary，仍需核对来源、作者许可及商标使用；3 份 notice 类文件要人工核查是否完整、是否随目标发行物提供。Rat 的“unapproved”是规则未批准，不是法律上认定侵权；同样，AL 也不验证作者是否真正有权许可。')
    add_p(d,'pom.xml 显式列出的 JUnit Jupiter、Commons IO、Mockito 均标为 test scope，但未解析父 POM、构建插件及传递依赖，不能据此宣称“无第三方依赖风险”。本次未构建 JAR，也未执行功能测试或漏洞检测。若用于正式交付，应以实际发行物为边界补充依赖解析、SBOM、许可证/版权扫描、二进制溯源及人工审批。')
    add_p(d,'本实验的充分性标准是：固定来源与工具版本、覆盖 116/116 个输入、可由 TXT/XML 独立核对统计、逐项说明 3 个待核问题。该标准满足本次限定的许可证头实验，不等于完整软件成分分析。',after=0)
    page_break(d)
    # Page 5
    heading(d,'4.4 Apache-2.0 使用义务与最终结论',2)
    add_p(d,'Apache-2.0 允许在满足条件时商用、修改及闭源再分发；分发时应提供许可证副本，对修改文件作显著标记，在分发的源代码中保留仍相关的版权、专利、商标及归属声明。如原作品含 NOTICE，应按第4(d)条在规定位置保留仍相关的归属内容。第3条规定特定专利诉讼会终止相应专利授权；第6条通常不授予商标使用权。不能只因“宽松许可证”就删掉 LICENSE/NOTICE。[2]')
    add_p(d,'结论：本次 116 个资源中，106 个标准文件识别为 Apache-2.0，3 个标准文件需人工复核，另有 3 个通知类和 4 个二进制文件需按类型核查。未在扫描范围内识别到另一种文本许可证，不等于不存在第三方许可问题。正式分发前须完成上页处置及实际交付物检查；本报告不作“全面合规”保证。')
    add_p(d,'实验执行与文字整理使用了 AI 辅助；提交前，学生应依据 README 亲自复现、理解 Unknown 的含义并复核结论。报告仅陈述已保存的实验事实，不虚构个人操作经历或学习感受。',after=5)
    heading(d,'依据与证据（访问/实验日期：2026-10-29）',2)
    link_p(d,'[1] Apache Commons CLI 的 GitHub 仓库：','https://github.com/apache/commons-cli')
    link_p(d,'[2] Apache License 2.0 原文：','https://www.apache.org/licenses/LICENSE-2.0')
    link_p(d,'[3] Apache Rat 官方说明：','https://creadur.apache.org/rat/')
    add_p(d,'[4] 固定源码、文件散列及原始输出：evidence/provenance.json、source-file-manifest.csv、rat-report.txt、rat-report.xml、source-excerpts/；命令与复现说明见 README.md。',indent=False,after=5)
    add_p(d,'选型同时参考教师提供的《33款开源许可合规工具—逐条优缺点与选型手册》；具体结论以固定源码、实测结果与许可证原文为依据。',indent=False,after=5)
    heading(d,'评价表格（保留教师原模板，待教师填写）',2)
    d._element.body.insert(len(d._element.body)-1,eval_xml)
    fix_table(d.tables[-1],size=10.5)
    for row in d.tables[-1].rows[1:]: row.cells[-1].text=''
    fix_table(d.tables[-1],size=10.5)
    path=ROOT/(STEM+'.docx'); d.save(path)
    # Searchable Markdown counterpart: same paragraphs/tables, ordered as Word.
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    md=['# '+TITLE,'']
    table_index=0
    for e in d._element.body:
        if e.tag==qn('w:p'):
            p=Paragraph(e,d)
            txt=p.text
            if txt: md.extend([txt,''])
        elif e.tag==qn('w:tbl'):
            t=Table(e,d)
            table_index+=1
            if table_index==1:
                md.extend(['| 字段 | 内容 |','| --- | --- |'])
                fields=[('姓名','宋佳威'),('学号','F2505437'),('班级','F250117'),('指导教师','祝欣蓉'),('开设学期','2026-2027秋季学期'),('开设时间','第1周——第16周'),('报告日期','2026年10月29日'),('评定成绩',''),('评定人','祝欣蓉'),('评定日期','')]
                md.extend('| '+key+' | '+value+' |' for key,value in fields)
                md.append('')
                continue
            for ix,row in enumerate(t.rows):
                md.append('| '+' | '.join(c.text.replace('\n',' / ') for c in row.cells)+' |')
                if ix==0: md.append('| '+' | '.join(['---']*len(row.cells))+' |')
            md.append('')
    # python-docx exposes hyperlink text via Paragraph.text in current version.
    (ROOT/(STEM+'.md')).write_text(chr(10).join(md)+chr(10),encoding='utf-8')
    if export_pdf:
        soffice=shutil.which('soffice')
        if not soffice: raise RuntimeError('DOCX generated; PDF export requires soffice')
        with tempfile.TemporaryDirectory(prefix='songjiawei-report-lo-') as td:
            profile=Path(td,'profile').as_uri()
            env=os.environ.copy()
            if Path('/System/Library/Fonts/Supplemental/Songti.ttc').exists():
                config=Path(td,'fonts.conf')
                config.write_text('<fontconfig><dir>/System/Library/Fonts/Supplemental</dir><cachedir>'+td+'/fontcache</cachedir><alias><family>宋体</family><prefer><family>Songti SC</family></prefer></alias><alias><family>SimSun</family><prefer><family>Songti SC</family></prefer></alias></fontconfig>')
                env['FONTCONFIG_FILE']=str(config)
            result=subprocess.run([soffice,'-env:UserInstallation='+profile,'--headless','--convert-to','pdf','--outdir',str(ROOT),str(path)],check=True,text=True,capture_output=True,timeout=120,env=env)
            print(result.stdout.strip())
        from pypdf import PdfReader
        reader=PdfReader(ROOT/(STEM+'.pdf'))
        print('PDF pages:',len(reader.pages))
        for i,p in enumerate(reader.pages,1):
            txt=p.extract_text() or ''
            print('page',i,'characters',len(txt),'starts',txt[:50].replace(chr(10),' '))
        assert len(reader.pages)==5,'Expected exactly five pages including cover and rubric'
        all_text=''.join(p.extract_text() or '' for p in reader.pages)
        for required in ['上海杉达学院','宋佳威','祝欣蓉','F2505437','F250117','开源项目介绍','工具输出的检测报告','检测结果详细分析','考核标准']:
            assert required in all_text, 'Missing or corrupt rendered text: '+required
    print(path)

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--template',type=Path,required=True);ap.add_argument('--no-pdf',action='store_true');a=ap.parse_args()
    build(a.template,not a.no_pdf)
